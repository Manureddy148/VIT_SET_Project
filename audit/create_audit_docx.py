"""One-shot script to generate audit/audit_1.docx."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from docx import Document
from docx.shared import Pt, RGBColor, Inches

doc = Document()

# ── Title ──────────────────────────────────────────────────────────────────────
title = doc.add_heading("Medical AI Pipeline — Audit Report", 0)
p = doc.add_paragraph()
p.add_run("Project: VIT_SET_Project-1  |  Date: April 7, 2026  |  Branch: main").bold = True
doc.add_paragraph()

# ── Layer table helper ─────────────────────────────────────────────────────────
STATUS_COLOR = {
    "BUILT":   RGBColor(0, 140, 0),
    "PARTIAL": RGBColor(180, 100, 0),
    "MISSING": RGBColor(200, 0, 0),
}


def add_layer_table(heading, items):
    doc.add_heading(heading, level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["Component (Docs)", "Status", "File / Notes"]):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
    for comp, status, note in items:
        row = t.add_row().cells
        row[0].text = comp
        row[2].text = note
        p2 = row[1].paragraphs[0]
        run = p2.add_run(status)
        run.bold = True
        run.font.color.rgb = STATUS_COLOR.get(status, RGBColor(50, 50, 50))
    doc.add_paragraph()


# ── Layers ────────────────────────────────────────────────────────────────────
doc.add_heading("PIPELINE vs CODE — Layer-by-Layer Audit", level=1)

add_layer_table("Layer 1 · Multimodal Input Router", [
    ("MultimodalRouter / detect_modality",       "BUILT",   "src/preprocessing/multimodal_router.py"),
    ("DICOMReader (pydicom optional import)",     "PARTIAL", "multimodal_router.py — optional try/except"),
    ("ECGLoader (CSV/EDF/XML text parse)",        "BUILT",   "multimodal_router.py"),
    ("AudioLoader (WAV/MP3)",                    "BUILT",   "multimodal_router.py"),
    ("Genomics upload handler (VCF/FASTQ)",      "BUILT",   "multimodal_router.py"),
    ("NiftiReader (nibabel — NIfTI/brain MRI)",  "MISSING", "—   nibabel not wired"),
    ("MissingDataAgent (dialogue flow)",         "PARTIAL", "safety_rails.py — prompts only"),
    ("FHIR R4 / HL7 format parsing",             "MISSING", "—"),
    ("FormatValidator (schema validation)",      "MISSING", "detect_modality() is filename heuristic only"),
])

add_layer_table("Layer 2 · Clinical Entity Extraction (NER)", [
    ("Bio_ClinicalBERT NER",                     "PARTIAL", "clinical_ner.py — gated MEDICAL_AI_USE_BIOMED_NER=1"),
    ("NegEx negation detection",                 "PARTIAL", "_is_negated() — 4-token window (not full NegEx)"),
    ("DomainClassifier",                         "BUILT",   "src/preprocessing/domain_classifier.py"),
    ("ECGProcessor (NeuroKit2)",                 "BUILT",   "src/preprocessing/ecg_processor.py"),
    ("ImagePreprocessor",                        "BUILT",   "multimodal_router.py — pixel mean/std stats"),
    ("FeatureNormalizer",                        "BUILT",   "src/preprocessing/normalizer.py"),
    ("UnitConverter",                            "MISSING", "—"),
    ("MissingValueImputer",                      "BUILT",   "FEATURE_MEDIANS fallback in each model.preprocess()"),
])

add_layer_table("Layer 3 · Model Registry (14 models in docs)", [
    ("DiabetesModel (XGBoost)",                  "BUILT",   "src/models/diabetes_model.py"),
    ("HeartModel (XGBoost)",                     "BUILT",   "src/models/heart_model.py"),
    ("PneumoniaModel",                           "PARTIAL", "src/models/pneumonia_model.py — rule-based stub"),
    ("CKDModel (LightGBM)",                      "MISSING", "—   not yet created"),
    ("SepsisModel (LightGBM)",                   "MISSING", "—   not yet created"),
    ("ImageRouter (ViT/CLIP)",                   "MISSING", "imaging_extension.py only — not wired"),
    ("ECGModel (1D-ResNet PTB-XL)",              "MISSING", "signal proxy used"),
    ("EchoModel (EchoNet)",                      "MISSING", "—"),
    ("BiomedCLIP zero-shot",                     "MISSING", "referenced in imaging_extension.py only"),
    ("xray_pneumonia ViT HF",                    "MISSING", "nickmuchi/vit model not loaded"),
    ("mri_brain EfficientNet HF",                "MISSING", "Devarshi/Brain_Tumor model not loaded"),
    ("skin_lesion HF",                           "MISSING", "—"),
    ("lung_sounds CNN HF",                       "MISSING", "audio pixel-proxy used"),
    ("polyp_endo HF",                            "MISSING", "—"),
    ("ModelRegistry",                            "BUILT",   "src/models/registry.py"),
])

add_layer_table("Layer 4 · SHAP + Ensemble", [
    ("SHAPExplainer (TreeExplainer)",            "BUILT",   "src/inference/shap_explainer.py"),
    ("SHAP → RAGQueryBuilder",                   "BUILT",   "src/rag/retriever.py — shap_to_rag_query()"),
    ("EnsembleVoter",                            "MISSING", "src/inference/ensemble.py — passthrough_score() stub"),
    ("KernelExplainer (NN)",                     "MISSING", "—"),
    ("PlattCalibrator",                          "MISSING", "—"),
    ("SeverityCalibrator",                       "PARTIAL", "domain amplifiers inside each model.predict()"),
    ("WaterfallPlotter",                         "MISSING", "SHAP bar chart in Streamlit only"),
])

add_layer_table("Layer 5 · RAG Pipeline", [
    ("ChromaDB VectorStore",                     "BUILT",   "src/rag/vector_store.py — PersistentClient + cosine"),
    ("PubMed Ingestor (Entrez)",                 "BUILT",   "scripts/ingest_pubmed.py — 73 real abstracts"),
    ("DomainFilter",                             "BUILT",   "vector_store.query(domain_filter=...)"),
    ("BioLORD-2023 Embeddings",                  "MISSING", "uses all-MiniLM-L6-v2 instead of FremyCompany/BioLORD-2023"),
    ("MMR Retriever (diversity=0.3)",            "MISSING", "simple cosine query — no Maximal Marginal Relevance"),
    ("HallucinationGuard (RAGAS)",               "PARTIAL", "snippet_guard by default; RAGAS gated MEDICAL_AI_USE_RAGAS=1"),
    ("Abstract coverage (500+/domain target)",   "PARTIAL", "73 total vs ~1500+ target"),
])

add_layer_table("Layer 6 · LLM Synthesis", [
    ("Llama-3-8B via Groq",                      "PARTIAL", "src/synthesis/llm_reporter.py — needs MEDICAL_AI_USE_LLM=1 + GROQ_API_KEY"),
    ("LangChain Prompt Template",                "PARTIAL", "llm_reporter.py — opt-in only"),
    ("SafetyRails",                              "BUILT",   "src/synthesis/safety_rails.py"),
    ("CriticalFlagger",                          "BUILT",   "build_safety_flags() → urgent_review flag"),
    ("FaithfulnessFilter",                       "BUILT",   "hallucination_guard.py — basic snippet check"),
    ("DisclaimerAppender",                       "BUILT",   "MEDICAL_DISCLAIMER constant appended"),
    ("StreamingReporter (WebSocket)",            "MISSING", "no WebSocket endpoint in main.py"),
])

add_layer_table("Layer 7 · Serving", [
    ("FastAPI (uvicorn)",                        "BUILT",   "src/api/main.py — 8+ endpoints"),
    ("Pydantic Schemas",                         "BUILT",   "src/api/schemas.py"),
    ("Streamlit UI",                             "BUILT",   "frontend/app.py — gauge + SHAP + citations + safety tabs"),
    ("JSONL Audit Log",                          "BUILT",   "src/api/audit.py — UUID + timestamp per prediction"),
    ("Docker + docker-compose",                  "BUILT",   "Dockerfile + docker-compose.yml present"),
    ("CORS Middleware",                          "MISSING", "not added to FastAPI app in main.py"),
    ("Authentication",                           "MISSING", "—   no API key / JWT"),
    ("WebSocket Stream",                         "MISSING", "—"),
    ("PDF Export (fpdf2)",                       "MISSING", "—   fpdf2 not installed/wired"),
    ("GCP Cloud Run config",                     "MISSING", "no cloudbuild.yaml or Cloud Run YAML"),
])

# ── Summary table ─────────────────────────────────────────────────────────────
doc.add_heading("Summary — Component Coverage per Layer", level=1)
t = doc.add_table(rows=1, cols=5)
t.style = "Table Grid"
for i, h in enumerate(["Layer", "Doc\nComponents", "Built ✓", "Partial ⚠", "Missing ✗"]):
    c = t.rows[0].cells[i]
    c.text = ""
    c.paragraphs[0].add_run(h).bold = True

for row_data in [
    ("L1 Input Router",   "9",  "4", "2", "3"),
    ("L2 NER/Preprocess", "8",  "5", "2", "1"),
    ("L3 Model Registry", "15", "3", "1", "11"),
    ("L4 SHAP/Ensemble",  "7",  "2", "1", "4"),
    ("L5 RAG Pipeline",   "7",  "3", "2", "2"),
    ("L6 LLM Synthesis",  "7",  "4", "2", "1"),
    ("L7 Serving",        "10", "5", "0", "5"),
    ("TOTAL",             "63", "26", "10", "27"),
]:
    r = t.add_row().cells
    for i, v in enumerate(row_data):
        r[i].text = v
        if row_data[0] == "TOTAL":
            r[i].paragraphs[0].runs[0].bold = True

doc.add_paragraph()

# ── Team guide checklist ──────────────────────────────────────────────────────
doc.add_heading("Team Guide Checklist Status", level=1)
t2 = doc.add_table(rows=1, cols=2)
t2.style = "Table Grid"
for i, h in enumerate(["Checklist Item", "Status"]):
    t2.rows[0].cells[i].text = ""
    t2.rows[0].cells[i].paragraphs[0].add_run(h).bold = True

checklist = [
    ("GitHub repo + folder structure",              "DONE",    RGBColor(0, 140, 0)),
    ("Python env + requirements.txt",               "DONE",    RGBColor(0, 140, 0)),
    ("Smoke test (medical_ai_e2e.py) — 8/8 PASS",  "DONE",    RGBColor(0, 140, 0)),
    ("DiabetesModel trained (AUC > 0.85)",          "DONE",    RGBColor(0, 140, 0)),
    ("HeartModel trained (AUC > 0.85)",             "DONE",    RGBColor(0, 140, 0)),
    ("SHAP explainer working",                      "DONE",    RGBColor(0, 140, 0)),
    ("FastAPI /predict returning JSON",             "DONE",    RGBColor(0, 140, 0)),
    ("Streamlit UI: gauge + SHAP chart",            "DONE",    RGBColor(0, 140, 0)),
    ("Safety disclaimer on all outputs",            "DONE",    RGBColor(0, 140, 0)),
    ("Docker container builds",                     "DONE",    RGBColor(0, 140, 0)),
    ("ingest_pubmed.py working (73 abstracts)",     "DONE",    RGBColor(0, 140, 0)),
    ("PneumoniaModel trained XGBoost",              "PARTIAL — rule-based only", RGBColor(180, 100, 0)),
    ("500+ abstracts/domain in ChromaDB",           "PARTIAL — 73 total",       RGBColor(180, 100, 0)),
    ("RAGAS faithfulness > 0.8",                    "PARTIAL — snippet guard",  RGBColor(180, 100, 0)),
    ("CKD + Sepsis models",                         "MISSING", RGBColor(200, 0, 0)),
    ("HuggingFace imaging models wired",            "MISSING", RGBColor(200, 0, 0)),
    ("BioLORD-2023 embeddings",                     "MISSING — using MiniLM",   RGBColor(200, 0, 0)),
    ("MMR retriever",                               "MISSING", RGBColor(200, 0, 0)),
    ("CORS + Auth on API",                          "MISSING", RGBColor(200, 0, 0)),
    ("PDF export (fpdf2)",                          "MISSING", RGBColor(200, 0, 0)),
    ("Deployed Streamlit Cloud / Render",           "MISSING", RGBColor(200, 0, 0)),
]

for item, status, color in checklist:
    r = t2.add_row().cells
    r[0].text = item
    run = r[1].paragraphs[0].add_run(status)
    run.bold = True
    run.font.color.rgb = color

doc.add_paragraph()

# ── Priority gaps ─────────────────────────────────────────────────────────────
doc.add_heading("Priority Implementation Gaps (Ordered)", level=1)
gaps = [
    "1. BioLORD-2023 replacing MiniLM in vector_store.py — high impact, doc-required",
    "2. MMR diversity in vector_store.query() — doc-required for retrieval quality",
    "3. Abstract volume — 73 vs 1500+ target (run ingest at higher --limit)",
    "4. CORS middleware — needed for any frontend/browser client to call API",
    "5. CKD + Sepsis models — Phase 2 checklist items",
    "6. PneumoniaModel — replace rule-based with trained XGBoost (Kaggle Chest X-Ray)",
    "7. HuggingFace imaging models (ViT xray, EfficientNet brain) — Phase 3",
    "8. WebSocket streaming endpoint — Phase 5",
    "9. PDF export (fpdf2) — Phase 5",
    "10. Deployment (Streamlit Cloud + Render) — Phase 5",
]
for g in gaps:
    doc.add_paragraph(g, style="List Number")

out_path = os.path.join(os.path.dirname(__file__), "audit_1.docx")
doc.save(out_path)
print(f"Saved: {out_path}")
