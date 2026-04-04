"""
Build docs/Medical_AI_Project_README_v1.docx — teammate handout (Word).

Requires: pip install python-docx
Run: python scripts/build_teammate_docx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "Medical_AI_Project_README_v1.docx"


def main() -> None:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        print("Install: python -m pip install python-docx")
        sys.exit(1)

    DOCS.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Medical AI Severity System", 0)
    doc.add_paragraph("README v1 — Teammate handout")
    doc.add_paragraph(
        "Research and academic project only. Not a medical device; not for diagnosis or treatment."
    )

    doc.add_heading("1. What this project is", level=1)
    doc.add_paragraph(
        "A demo of a seven-layer pipeline: clinical text and structured labs → domain routing "
        "(diabetes / heart disease / pneumonia) → ML severity score 0–100 with feature attribution "
        "(SHAP) → optional literature retrieval (ChromaDB) → written explanation with a fixed safety disclaimer → "
        "FastAPI backend and Streamlit UI."
    )

    doc.add_heading("2. Main folders", level=1)
    for line in [
        "src/api — FastAPI: GET /health, POST /predict",
        "src/models — XGBoost models + registry + pneumonia rule stub",
        "src/preprocessing — light lab extraction from text",
        "src/rag — vector store (optional if MEDICAL_AI_SKIP_RAG=1)",
        "src/synthesis — report text + disclaimer",
        "frontend/ — Streamlit dashboard",
        "tests/ — pytest, including test_e2e_smoke.py",
        "launch_api.py — start API and open Swagger",
    ]:
        p = doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("3. How to run (Windows)", level=1)
    doc.add_paragraph("Install: python -m pip install -r requirements.txt")
    doc.add_paragraph("Tests: pytest")
    doc.add_paragraph("API (keep terminal open): python launch_api.py → http://127.0.0.1:8000/docs")
    doc.add_paragraph("UI: streamlit run frontend/app.py (API must be running).")

    doc.add_heading("4. Environment variables", level=1)
    doc.add_paragraph(
        "MEDICAL_AI_SKIP_RAG=1 — skip Chroma (default in tests). "
        "MEDICAL_API_URL — Streamlit API base. MODEL_REGISTRY_DIR — saved .joblib path."
    )

    doc.add_heading("5. Models", level=1)
    doc.add_paragraph(
        "Without Kaggle exports, the API trains small synthetic XGBoost models at startup. "
        "For saved weights: python scripts/train_models.py"
    )

    doc.add_heading("6. Full documentation", level=1)
    doc.add_paragraph(
        "See README.md (GitHub root), medical-ai-project-guide.md, QUICKSTART.txt, and "
        "Create Roadmap/Medical_AI_Team_Guide.html for the full semester plan and architecture."
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
